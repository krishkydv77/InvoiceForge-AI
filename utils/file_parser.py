import csv
import os

import openpyxl
import pdfplumber
import pytesseract
from PIL import Image
from docx import Document


def parse_pdf(path: str) -> str:
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                chunks.append(page_text)
    text = "\n".join(chunks).strip()
    if text:
        return text

    # Scanned/image-only PDF fallback.
    try:
        from pdf2image import convert_from_path
    except ImportError as exc:
        raise RuntimeError(
            "This PDF contains no extractable text. Install pdf2image and Poppler "
            "to enable OCR fallback for scanned PDFs."
        ) from exc

    lang = os.getenv("OCR_LANG", "eng").strip() or "eng"
    pages = convert_from_path(path)
    return "\n".join(pytesseract.image_to_string(page, lang=lang) for page in pages).strip()


def parse_image(path: str, lang: str | None = None) -> str:
    lang = lang or os.getenv("OCR_LANG", "eng").strip() or "eng"
    try:
        with Image.open(path) as img:
            text = pytesseract.image_to_string(img, lang=lang)
    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError(
            "Tesseract OCR is not installed/configured on this machine. "
            "Install Tesseract and add it to PATH."
        ) from exc
    return text.strip()


def parse_docx(path: str) -> str:
    doc = Document(path)
    chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(chunks).strip()


def parse_excel(path: str) -> str:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(" | ".join(str(cell) if cell is not None else "" for cell in row))
        return "\n".join(rows).strip()
    finally:
        wb.close()


def parse_csv(path: str) -> str:
    chunks = []
    with open(path, newline="", encoding="utf-8-sig") as csvfile:
        reader = csv.reader(csvfile)
        for row in reader:
            chunks.append(" | ".join(row))
    return "\n".join(chunks).strip()
